"""
Simple GUI tool to convert newline-separated strings into a PDF
containing QR codes with captions.
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox

import qrcode
from openpyxl import load_workbook
from PIL import Image
from reportlab.lib import pagesizes, units
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


PAGE_SIZE = pagesizes.A6
MARGIN = 12 * units.mm
QR_SIZE = 65 * units.mm
CAPTION_GAP = 8


def generate_qr_image(data: str, pixel_size: int = 512) -> Image.Image:
    """Return a PIL image of the QR code for the provided data string."""
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=4,
    )
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white").convert("RGB")
    return img.resize((pixel_size, pixel_size), Image.LANCZOS)


def add_qr_to_canvas(pdf: canvas.Canvas, data: str) -> None:
    """Draw a QR code centered on the page with a caption underneath."""
    qr_image = generate_qr_image(data)
    qr_bytes = io.BytesIO()
    qr_image.save(qr_bytes, format="PNG")
    qr_bytes.seek(0)

    page_width, page_height = PAGE_SIZE
    col_x = (page_width - QR_SIZE) / 2
    qr_bottom = page_height - MARGIN - QR_SIZE

    pdf.drawImage(
        ImageReader(qr_bytes),
        col_x,
        qr_bottom,
        width=QR_SIZE,
        height=QR_SIZE,
        preserveAspectRatio=True,
        mask="auto",
    )

    caption_y = qr_bottom - CAPTION_GAP
    pdf.setFont("Helvetica", 9)
    pdf.drawCentredString(page_width / 2, caption_y, data)


def build_pdf(entries: list[str], output_path: str) -> None:
    """Create a PDF with QR codes for every entry (one per page)."""
    pdf = canvas.Canvas(output_path, pagesize=PAGE_SIZE)

    for data in entries:
        add_qr_to_canvas(pdf, data)
        pdf.showPage()

    pdf.save()


class QRGui(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("QR PDF Generator")
        self.geometry("600x420")
        self.resizable(True, True)

        self._build_widgets()

    def _build_widgets(self) -> None:
        instructions = (
            "Вставьте список строк (по одной на каждой строке). "
            "Для каждой строки будет создан отдельный QR-код."
        )
        tk.Label(self, text=instructions, wraplength=560, justify="left").pack(
            pady=(10, 5)
        )

        self.text_box = tk.Text(self, height=15, wrap="none", font=("Consolas", 11))
        self._bind_shortcuts()
        self.text_box.pack(fill="both", expand=True, padx=10, pady=5)

        button_frame = tk.Frame(self)
        button_frame.pack(pady=10)

        tk.Button(
            button_frame, text="Сгенерировать PDF", command=self.generate_pdf
        ).pack(side="left", padx=5)
        tk.Button(button_frame, text="Загрузить Excel", command=self.load_excel).pack(
            side="left", padx=5
        )
        tk.Button(button_frame, text="Очистить", command=self.clear_text).pack(
            side="left", padx=5
        )
        tk.Button(button_frame, text="Выход", command=self.destroy).pack(
            side="left", padx=5
        )

    def clear_text(self) -> None:
        self.text_box.delete("1.0", tk.END)

    def generate_pdf(self) -> None:
        raw_text = self.text_box.get("1.0", tk.END).strip()
        entries = [line.strip() for line in raw_text.splitlines() if line.strip()]

        if not entries:
            messagebox.showwarning("Пустой список", "Введите хотя бы одну строку.")
            return

        folder = filedialog.askdirectory(title="Выберите папку для сохранения PDF")

        if not folder:
            return

        output_path = self._build_output_path(folder)

        try:
            build_pdf(entries, output_path)
        except Exception as exc:  # pylint: disable=broad-except
            messagebox.showerror("Ошибка", f"Не удалось создать PDF: {exc}")
            return

        messagebox.showinfo("Готово", f"PDF файл сохранён:\n{output_path}")

    def load_excel(self) -> None:
        file_path = filedialog.askopenfilename(
            title="Выберите Excel файл",
            filetypes=[("Excel файлы", "*.xlsx;*.xlsm;*.xltx;*.xltm")],
        )

        if not file_path:
            return

        try:
            entries = self._read_excel_entries(file_path)
        except Exception as exc:  # pylint: disable=broad-except
            messagebox.showerror("Ошибка", f"Не удалось прочитать Excel: {exc}")
            return

        if not entries:
            messagebox.showwarning("Пустой файл", "В выбранном файле нет данных.")
            return

        self._insert_entries(entries)

    def _bind_shortcuts(self) -> None:
        def handle_ctrl_key(event: tk.Event) -> str | None:  # type: ignore[name-defined]
            keysym = (event.keysym or "").lower()
            if keysym in {"v", "в"}:
                self.text_box.event_generate("<<Paste>>")
                return "break"
            return None

        self.text_box.bind("<Control-KeyPress>", handle_ctrl_key, add="+")

    def _insert_entries(self, entries: list[str]) -> None:
        content = "\n".join(entries)
        if self.text_box.get("1.0", tk.END).strip():
            self.text_box.insert(tk.END, "\n" + content)
        else:
            self.text_box.insert("1.0", content)

    @staticmethod
    def _read_excel_entries(file_path: str) -> list[str]:
        workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
        sheet = workbook.active
        entries: list[str] = []

        for row in sheet.iter_rows(values_only=True):
            if not row:
                continue
            value = row[0]
            if value is None:
                continue
            entry = str(value).strip()
            if entry:
                entries.append(entry)

        return entries

    @staticmethod
    def _build_output_path(folder: str) -> str:
        today = date.today().strftime("%Y-%m-%d")
        base_name = f"Коды_коробов_{today}"
        target_folder = Path(folder)

        counter = 1
        while True:
            candidate = target_folder / f"{base_name}_{counter}.pdf"
            if not candidate.exists():
                return str(candidate)
            counter += 1


def main() -> None:
    app = QRGui()
    app.mainloop()


if __name__ == "__main__":
    main()

